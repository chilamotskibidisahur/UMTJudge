import docx

doc = docx.Document('De_Cuong_KHKT_ChauCaoJudge_v2.docx')
with open('docx_dump.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            style_name = p.style.name if p.style else ''
            f.write(f'P{i} [{style_name}]: {p.text}\n')
    for ti, t in enumerate(doc.tables):
        f.write(f'\n=== TABLE {ti} ===\n')
        for row in t.rows:
            f.write(str([c.text.strip().replace('\n', ' ') for c in row.cells]) + '\n')
print('Successfully dumped to docx_dump.txt')
