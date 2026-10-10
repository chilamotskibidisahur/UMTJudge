# -*- coding: utf-8 -*-
"""So sánh trùng lặp >= 6 từ liên tiếp giữa bản đề cương mới và file mẫu."""
import glob, re, sys, unicodedata, json
import docx

N = 6


def texts(path):
    d = docx.Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                out.append(c.text)
    return [x for x in out if x.strip()]


def words(s):
    s = unicodedata.normalize("NFC", s.lower())
    return re.findall(r"\w+", s)


def grams(ws):
    return {" ".join(ws[i:i + N]) for i in range(len(ws) - N + 1)}


tpl = [f for f in glob.glob("*.docx") if "KHKT" in f and "ChauCao" not in f][0]
tpl_grams = set()
for t in texts(tpl):
    tpl_grams |= grams(words(t))

hits = []
for t in texts("output/De_Cuong_KHKT_ChauCaoJudge.docx"):
    g = grams(words(t))
    common = sorted(g & tpl_grams)
    if common:
        hits.append({"paragraph": t[:120], "matches": common})

res = {"template": tpl, "ngram": N, "paragraphs_with_overlap": len(hits), "hits": hits}
sys.stdout.buffer.write(json.dumps(res, ensure_ascii=False, indent=1).encode("utf-8"))
